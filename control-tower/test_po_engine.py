"""
PO Automation as a TRANSACTION ENGINE — the invariants, the failure matrix,
crash consistency, concurrency and properties.

    python test_po_engine.py

Everything here runs against STAND-INS (fixtures/po_engine.py: a mocked eHub
source, a stand-in Microsoft Graph over HTTP). It proves the engine's rules —
it is NOT evidence that the real eHub, the real Bill of Entry PDFs or the
real Graph mailbox behave this way. That needs the real Windows worker
(PO_WORKER_RUNBOOK.md).
"""

import hashlib
import json
import os
import random
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE / "fixtures"))
WORK = Path(tempfile.mkdtemp(prefix="po_engine_"))
os.environ["PO_DATA_DIR"] = str(WORK / "default")
os.environ["PO_ALLOW_TEST_SEND"] = "1"

import po_engine as F  # noqa: E402
from po import doctypes, extract as X, mail as M, pipeline as P, quality as Q  # noqa: E402
from po import store as S, template as T, validate as V  # noqa: E402

GRAPH = F.GraphStub()
os.environ.update(GRAPH.env())
CONFIG = {"recipient": "accounts.ghana@mantrac.test", "sender": "ata@mantrac.test",
          "auto_send": False, "defaults": {"supplier": "CAT", "branch": None, "charge_to": None,
                                           "priority": None}, "confirm_wait_s": 0}
NOSLEEP = lambda s: None  # noqa: E731
PASS, FAIL = [], []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print("  {0}  {1}{2}".format("PASS" if ok else "FAIL", name,
                                 "  ({0})".format(str(detail)[:400]) if detail and not ok else ""))


def rule(title):
    print("\n" + "=" * 74 + "\n" + title + "\n" + "=" * 74)


_n = {"i": 0}


def fresh_ref():
    _n["i"] += 1
    return "176-{0:08d}".format(70000000 + _n["i"])


def store(name):
    return S.Store(folder=WORK / name)


def job(st, ref=None, data=None, request=None, source=None, **text):
    ref = ref or fresh_ref()
    if data is None:
        data = F.pdf_of(F.boe_text(bl=ref, **text))
    record = st.create(doctypes.DEFAULT, ref, request or {})
    return P.process(st, record, source or F.Source(data, ref=ref), CONFIG, sleep=NOSLEEP)


def sends_for(record):
    subject = (record.get("email") or {}).get("subject")
    return GRAPH.sent_for(subject=subject) if subject else []


def invariants(st, record):
    """The engine's invariants for one job — [] when every one holds."""
    bad = []
    ms = [m["name"] for m in record.get("milestones") or []]
    have = set(ms)
    for name in ms:
        missing = [p for p in S.PREREQUISITES.get(name, ()) if p not in have]
        if missing:
            bad.append("{0} recorded without {1}".format(name, missing))
    v = record.get("validation") or {}
    state = record["state"]
    after_valid = (S.VALIDATED, S.TEMPLATE_GENERATED, S.SAVED, S.EMAIL_PREPARED, S.EMAIL_SENDING,
                   S.EMAIL_SENT, S.EMAIL_CONFIRMED, S.EMAIL_UNKNOWN,
                   S.EMAIL_RECONCILIATION_FAILED)
    if state in after_valid and v.get("decision") != "VALID":
        bad.append("{0} with validation {1}".format(state, v.get("decision")))
    if v and v.get("decision") != "VALID" and (sends_for(record) or record.get("output")):
        bad.append("an output or email exists for a job whose validation is not VALID")
    if (record.get("identity") or {}).get("decision") == "MISMATCH" and state in after_valid:
        bad.append("identity MISMATCH yet {0}".format(state))
    if "FIELDS_EXTRACTED" in have and "PDF_INTEGRITY_VERIFIED" not in have:
        bad.append("fields extracted without PDF integrity")
    if record.get("output"):
        g4 = (record.get("request_fields") or {}).get("invoice_no") or {}
        printed = record.get("invoice_candidate") or {}
        allowed = set(([printed.get("value")] if printed.get("value") else []) +
                      list(printed.get("candidates") or []))
        if not str(g4.get("origin") or "").startswith("bill_of_entry") or \
                g4.get("value") not in allowed:
            bad.append("G4 {0!r} is not an explicit printed Invoice No. ({1})".format(
                g4.get("value"), g4.get("origin")))
    if state == S.EMAIL_CONFIRMED:
        out = record.get("output") or {}
        if not out.get("path") or not Path(out["path"]).is_file() or \
                hashlib.sha256(Path(out["path"]).read_bytes()).hexdigest() != out.get("sha256"):
            bad.append("confirmed, but the output is missing or changed")
        if not ((record.get("email") or {}).get("reconciliation") or {}).get("verified"):
            bad.append("confirmed without a verified reconciliation")
        if "COMPLETED" not in have:
            bad.append("confirmed without COMPLETED")
    if len(sends_for(record)) > 1:
        bad.append("{0} emails for one document".format(len(sends_for(record))))
    audit = st.verify_audit()
    if not audit["ok"]:
        bad.append("audit chain broken: " + str(audit["problem"]))
    return bad


# ═════════════════════════════════════════════════════════════════════════
rule("1. THE CANONICAL TRANSACTION — every milestone, in order, with evidence")
# ═════════════════════════════════════════════════════════════════════════
st = store("canon")
r = job(st)
r, out = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=NOSLEEP)
names = [m["name"] for m in r["milestones"]]
check("A good job, sent: EMAIL_CONFIRMED, one email, recipient/subject/attachment verified",
      out == "CONFIRMED" and r["state"] == S.EMAIL_CONFIRMED and len(sends_for(r)) == 1
      and r["email"]["reconciliation"]["verified"], (out, r["state"]))
check("All 21 canonical milestones are recorded, in the canonical order",
      names == list(S.CANONICAL), names)
check("Every milestone carries evidence (status observed, document hash, decision, …)",
      all(m.get("evidence") for m in r["milestones"]),
      [m["name"] for m in r["milestones"] if not m.get("evidence")])
ev = {m["name"]: m["evidence"] for m in r["milestones"]}
check("ELIGIBLE_VERIFIED records the observed status, its source and when",
      ev["ELIGIBLE_VERIFIED"]["observed_status"] == "Under Clearance"
      and ev["ELIGIBLE_VERIFIED"]["source"] == "eHub shipment list"
      and ev["ELIGIBLE_VERIFIED"]["at"], ev["ELIGIBLE_VERIFIED"])
check("PDF_RETRIEVED / PDF_INTEGRITY_VERIFIED record the hash, size and the checks",
      ev["PDF_RETRIEVED"]["sha256"] == r["document"]["sha256"]
      and ev["PDF_INTEGRITY_VERIFIED"]["eof_marker"] and ev["PDF_INTEGRITY_VERIFIED"]["verified"])
check("OUTPUT_PERSISTED records the absolute path, size, hash and job",
      Path(ev["OUTPUT_PERSISTED"]["path"]).is_absolute() and ev["OUTPUT_PERSISTED"]["sha256"]
      == r["output"]["sha256"] and ev["OUTPUT_PERSISTED"]["job"] == r["po_id"])
check("The job carries its transaction metadata: attempt, previous state, started/completed, "
      "worker, idempotency key",
      r["attempt"] == 1 and r["previous_state"] == S.EMAIL_SENT and r["started_at"]
      and r["completed_at"] and r["worker"]["pid"] and r["idempotency_key"], r.get("worker"))
check("Invariants hold", not invariants(st, r), invariants(st, r))

# ═════════════════════════════════════════════════════════════════════════
rule("2. THE STATE LAYER ENFORCES THE INVARIANTS — whatever code asks")
# ═════════════════════════════════════════════════════════════════════════
st = store("inv")
base = st.create(doctypes.DEFAULT, "176-11110000", {})


def refused(record, to, words):
    try:
        st.transition(record, to)
        return False
    except S.InvariantViolation as error:
        return all(w in str(error) for w in words)
    except S.IllegalTransition:
        return False


check("FIELDS_EXTRACTED without PDF_INTEGRITY_VERIFIED: refused",
      refused(dict(base, state=S.PDF_READ, milestones=[]), S.FIELDS_EXTRACTED,
              ["FIELDS_EXTRACTED"]))
check("VALIDATED without BUSINESS_RULES_VALIDATED / a VALID decision: refused",
      refused(dict(base, state=S.VALIDATING), S.VALIDATED, ["BUSINESS_RULES_VALIDATED", "VALID"]))
check("TEMPLATE_GENERATED without TEMPLATE_GENERATED + TEMPLATE_VERIFIED milestones: refused",
      refused(dict(base, state=S.VALIDATED), S.TEMPLATE_GENERATED, ["TEMPLATE_VERIFIED"]))
check("EMAIL_PREPARED without the EMAIL_PREPARED milestone and a verified output: refused",
      refused(dict(base, state=S.SAVED), S.EMAIL_PREPARED, ["EMAIL_PREPARED", "output"]))
check("EMAIL_SENDING without IDEMPOTENCY_CONFIRMED: refused",
      refused(dict(base, state=S.EMAIL_PREPARED), S.EMAIL_SENDING, ["IDEMPOTENCY_CONFIRMED"]))
check("EMAIL_CONFIRMED without reconciliation evidence and COMPLETED: refused",
      refused(dict(base, state=S.EMAIL_SENT), S.EMAIL_CONFIRMED, ["COMPLETED", "reconciled"]))
try:
    st.milestone(dict(base), "TEMPLATE_GENERATED")
    ok = False
except S.InvariantViolation as error:
    ok = "BUSINESS_RULES_VALIDATED" in str(error)
check("A milestone before its prerequisites (TEMPLATE_GENERATED first): refused", ok)
sneaky = dict(base, state=S.VALIDATING, validation={"decision": "VALID", "passed": True},
              identity={"decision": "MATCH"},
              request_fields={"invoice_no": {"value": "70076", "origin": "request"}},
              milestones=[{"name": n} for n in S.CANONICAL[:12]])
check("VALIDATED with G4 from anywhere but the Bill of Entry (e.g. UNA+ 70076): refused",
      refused(sneaky, S.VALIDATED, ["G4"]))
check("COMPLETED is unreachable except through EMAIL_CONFIRMED",
      all(S.EMAIL_CONFIRMED in [x for x in S.PREREQUISITES["COMPLETED"]] for _ in [0]))

# ═════════════════════════════════════════════════════════════════════════
rule("3. DURABLE PERSISTENCE — versioned writes, atomic files, OS locks")
# ═════════════════════════════════════════════════════════════════════════
st = store("persist")
a = st.create(doctypes.DEFAULT, "176-22220000", {})
b = st.get(a["po_id"])
a["progress"] = "writer A"
st.save(a)
try:
    b["progress"] = "writer B (stale copy)"
    st.save(b)
    lost = True
except S.ConcurrentUpdate:
    lost = False
check("Two copies of one job: the second write is REFUSED (no lost update)",
      not lost and st.get(a["po_id"])["progress"] == "writer A")
check("No temporary file is left after writes",
      not list((st.folder / "jobs").glob("*.tmp")) and not list(st.folder.glob("*.tmp")))
child = subprocess.Popen([sys.executable, "-c", (
    "import sys,time; sys.path.insert(0,{0!r}); from po import store as S\n"
    "st=S.Store(folder={1!r})\n"
    "cm=st.xlock('ledger'); cm.__enter__(); print('held', flush=True); time.sleep(60)").format(
        str(HERE), str(st.folder))], stdout=subprocess.PIPE, text=True)
child.stdout.readline()
t0 = time.time()
try:
    with st.xlock("ledger", timeout=0.5):
        got_while_held = True
except TimeoutError:
    got_while_held = False
child.kill()
child.wait()
t1 = time.time()
with st.xlock("ledger", timeout=5):
    waited = time.time() - t1
check("An OS lock held by another process blocks; once that process is KILLED it is free at "
      "once (no stale lock to break)", not got_while_held and waited < 1.0,
      (got_while_held, round(waited, 2)))
(st.folder / "locks" / "claims.lock").write_text("99999 0")
with st.xlock("claims", timeout=1):
    stale_ok = True
check("A leftover lock file from a dead process does not block", stale_ok)

# ═════════════════════════════════════════════════════════════════════════
rule("4. AUDIT TRAIL — every transition recorded; append-only; tamper-evident")
# ═════════════════════════════════════════════════════════════════════════
st = store("audit")
r = job(st)
events = st.events(r["po_id"], limit=5000)
changes = [e for e in events if e["event"] == "STATE_CHANGED"]
check("Every state change is an event with old and new state, actor, attempt, correlation id",
      [c["metadata"]["new_state"] for c in changes] == [h["state"] for h in r["history"][1:]]
      and all(c["metadata"]["old_state"] and c["actor"] and c["attempt"] and
              c["correlation_id"] == r["po_id"] for c in changes),
      [c["metadata"]["new_state"] for c in changes])
check("The chain verifies", st.verify_audit()["ok"], st.verify_audit())
path = st.folder / "events.jsonl"
original = path.read_text(encoding="utf-8")
lines = original.splitlines()
path.write_text("\n".join(lines[:5] + [lines[5].replace("OK", "FAILED", 1)] + lines[6:]) + "\n",
                encoding="utf-8")
altered = st.verify_audit()
path.write_text("\n".join(lines[:5] + lines[6:]) + "\n", encoding="utf-8")
removed = st.verify_audit()
path.write_text("\n".join(lines[:5] + [lines[6], lines[5]] + lines[7:]) + "\n", encoding="utf-8")
reordered = st.verify_audit()
path.write_text("\n".join(lines[:-1]) + "\n", encoding="utf-8")
truncated = st.verify_audit()
path.write_text(original, encoding="utf-8")
check("An altered event is detected", not altered["ok"] and "altered" in altered["problem"],
      altered)
check("A removed event is detected", not removed["ok"], removed)
check("Reordered events are detected", not reordered["ok"], reordered)
check("Events cut from the end are detected (the head no longer matches)",
      not truncated["ok"], truncated)
check("...and the untouched trail verifies again", st.verify_audit()["ok"])
blob = original.lower()
check("No secret is stored in the trail (password, token, cookie, authorization, security code)",
      not any(w in blob for w in ('"password"', '"token"', '"cookie"', '"authorization"',
                                  "security_code", "client_secret")))

# ═════════════════════════════════════════════════════════════════════════
rule("5. NUMBERS — one grammar; malformed values never coerced")
# ═════════════════════════════════════════════════════════════════════════
accepted = {"1,234.56": 1234.56, "653,492.35": 653492.35, "1234": 1234.0, "0.5": 0.5,
            "11.2045": 11.2045, "-12.50": -12.5, "(12.50)": -12.5, "1,234,567.89": 1234567.89}
rejected = ("1,2,3", "12,34,567", "1.2.3", "10O0", "1O00", "1.234,56", "1234,56", "12,50",
            "1 234.56", ".5", "1,23", "12a", "", None)
check("Accepted forms read exactly", all(X.to_float(k) == v for k, v in accepted.items()),
      {k: X.to_float(k) for k in accepted})
check("Rejected forms are None — never a plausible number",
      all(X.to_float(k) is None for k in rejected), {k: X.to_float(k) for k in rejected})
for label, kw in (("total duty '653,4O2.35'", {"duty": "653,4O2.35"}),
                  ("exchange rate '11.2O'", {"rate": "11.2O"}),
                  ("import duty '246.320,05'", {"import_duty": "246.320,05"})):
    st = store("num" + label[:6].strip())
    r = job(st, **kw)
    check("Malformed {0}: EXTRACTION_FAILED (NORMALIZATION_FAILED), raw kept, nothing "
          "generated".format(label), r["state"] == S.EXTRACTION_FAILED and r["failure"]["code"]
          == "NORMALIZATION_FAILED" and not r.get("output") and not invariants(st, r),
          (r["state"], r.get("failure")))

# ═════════════════════════════════════════════════════════════════════════
rule("6. VAT / MULTI-ITEM — reconciled by the document's own rule, or stopped")
# ═════════════════════════════════════════════════════════════════════════
st = store("vat")
r = job(st)
calc = {c["rule"]: c for c in r["calculations"]}
check("One VAT block: Σ lines = 407,172.30; duty − VAT = the printed import duty (exact)",
      calc["G20_VAT_BLOCK"]["result"] == "407172.30"
      and calc["XCHECK_IMPORT_DUTY"]["result"] == "OK" and r["state"] == S.EMAIL_PREPARED)
r = job(st, vat="Import VAT                        304,446.44\n"
                "Network Charge VAT                      0.00\n"
                "Import NHIL                        50,741.08\n"
                "GETFund Import                     50,741.08\n"
                "Network Charge VAT Fund Levy        1,243.70\n")
check("A zero VAT line is kept (not dropped) and the block still reconciles",
      any(l["amount"] == 0.0 for l in r["fields"]["vat_lines"]["value"])
      and r["state"] == S.EMAIL_PREPARED, (r["state"], r["fields"]["vat_lines"].get("value")))
r = job(st, vat="Import VAT                        152,223.22\n"
                "Import VAT                        152,223.22\n"
                "Network Charge VAT                  1,066.03\n")
check("Multi-item: the same VAT label printed twice (a repeat, or a second item?) — the rule "
      "cannot tell, so AMBIGUOUS: stopped, never summed twice",
      r["fields"]["vat_lines"]["status"] == X.AMBIGUOUS and r["state"] == S.VALIDATION_FAILED
      and not r.get("output"), r["state"])
r = job(st, vat="")
check("No VAT block at all: MISSING — stopped, nothing assumed zero",
      r["fields"]["vat_lines"]["status"] == X.MISSING and r["state"] == S.VALIDATION_FAILED)
r = job(st, import_duty="246,318.80")
check("Rounding inside the 1.00 GHS tolerance (variance 1.25 → outside) is decided exactly",
      r["state"] == S.VALIDATION_FAILED and
      {c["rule"]: c for c in r["calculations"]}["XCHECK_IMPORT_DUTY"]["intermediate"] == "1.25",
      {c["rule"]: c for c in r["calculations"]}["XCHECK_IMPORT_DUTY"])
r = job(st, import_duty="246,319.10")
check("...a variance of exactly 0.95 is inside: VALID", r["state"] == S.EMAIL_PREPARED,
      r["state"])
r = job(st, duty="-653,492.35")
check("A negative (credit) total duty: read as negative, stopped by the > 0 rule",
      r["fields"]["duty_amount_ghs"]["value"] < 0 and r["state"] == S.VALIDATION_FAILED)
f1 = X.extract(F.boe_text(), doctypes.get())
check("Deterministic: the same inputs always give the same calculations",
      V.calculations(doctypes.get(), f1) == V.calculations(doctypes.get(), f1))

# ═════════════════════════════════════════════════════════════════════════
rule("7. G4 — only the Bill of Entry's explicit Invoice No.")
# ═════════════════════════════════════════════════════════════════════════
st = store("g4")
cases = [("explicit", {"invoice": "2600005261"}, {}, S.EMAIL_PREPARED, None),
         ("absent (UNA+ only)", {"invoice": None, "extra": "UNA+ Invoice Number 70076\n"}, {},
          S.NEEDS_REVIEW, "absent"),
         ("absent, value given with the job", {"invoice": None}, {"invoice_no": "70076"},
          S.NEEDS_REVIEW, "absent"),
         ("two printed", {"invoice": "A1001", "extra": "Commercial Invoice Number: B2002\n"},
          {}, S.NEEDS_REVIEW, "ambiguous"),
         ("malformed", {"invoice": "2600 005261"}, {}, S.NEEDS_REVIEW, "malformed"),
         ("the job disagrees", {"invoice": "2600005261"}, {"invoice_no": "999"},
          S.NEEDS_REVIEW, "conflict"),
         ("the job agrees", {"invoice": "2600005261"}, {"invoice_no": "2600005261"},
          S.EMAIL_PREPARED, None)]
for label, text, request, want, reason in cases:
    r = job(st, request=request, **text)
    check("G4 {0}: {1}{2}".format(label, want, " ({0})".format(reason) if reason else ""),
          r["state"] == want and (r.get("failure") or {}).get("g4_reason") == reason
          and not invariants(st, r) and ("70076" not in json.dumps(r.get("request_fields"))),
          (r["state"], r.get("failure"), invariants(st, r)))

# ═════════════════════════════════════════════════════════════════════════
rule("8. SHIPMENT IDENTITY — MATCH, MISMATCH, INSUFFICIENT_EVIDENCE")
# ═════════════════════════════════════════════════════════════════════════
st = store("ident")
ref = fresh_ref()
r = job(st, ref=ref, source=F.Source(F.pdf_of(F.boe_text(bl=ref)), ref=ref,
                                     page_shows="176-99999999", on_manage=False))
check("The Manage page shows another BOL/AWB: IDENTITY_MISMATCH before anything is read",
      r["state"] == S.IDENTITY_MISMATCH and not r.get("fields")
      and "PDF_INTEGRITY_VERIFIED" not in S.Store.milestones(r), r["state"])
ref = fresh_ref()
r = job(st, ref=ref, source=F.Source(F.pdf_of(F.boe_text(bl=ref)), ref=ref,
                                     page_declaration="40799999999"))
check("The Manage page's declaration is not the selected Bill Entry's: IDENTITY_MISMATCH",
      r["state"] == S.IDENTITY_MISMATCH, r["state"])
ref = fresh_ref()
r = job(st, ref=ref, data=F.pdf_of(F.boe_text(bl="176-12345678")))
check("The PDF's BL/AWB is another shipment's: IDENTITY_MISMATCH, nothing generated",
      r["state"] == S.IDENTITY_MISMATCH and r["identity"]["decision"] == "MISMATCH"
      and not r.get("output"), r["state"])
ref = fresh_ref()
r = job(st, ref=ref, source=F.Source(F.pdf_of(F.boe_text(bl=ref)), ref=ref, on_manage=False))
check("The Manage page does not show the BOL/AWB, but the PDF's BL/AWB and declaration tie the "
      "document to the row and eHub's Bill Entry: MATCH (proven by the document)",
      r["state"] == S.EMAIL_PREPARED and r["identity"]["decision"] == "MATCH"
      and r["identity"]["stage"] == "document", (r["state"], r.get("identity")))
ref = fresh_ref()
r = job(st, ref=ref, source=F.Source(F.pdf_of(F.boe_text(bl=ref)), ref=ref, on_manage=False,
                                     identifier=None))
check("Neither the page nor the document proves it: INSUFFICIENT_EVIDENCE → NEEDS_REVIEW "
      "(IDENTITY_UNPROVEN) — never treated as MATCH",
      r["state"] in (S.NEEDS_REVIEW, S.VALIDATION_FAILED) and
      r["identity"]["decision"] == "INSUFFICIENT_EVIDENCE" and not r.get("output"),
      (r["state"], r.get("identity"), (r.get("validation") or {}).get("code")))

# ═════════════════════════════════════════════════════════════════════════
rule("9. VALIDATION GATE — one decision, structured")
# ═════════════════════════════════════════════════════════════════════════
st = store("gate")
r = job(st)
v = r["validation"]
check("VALID: decision, checks, comparisons (MATCH per field), calculations, no failures",
      v["decision"] == "VALID" and v["comparisons"] and
      all(c["result"] in ("MATCH", "NOT_AVAILABLE") for c in v["comparisons"])
      and v["calculations"] and not v["failed_checks"], v["comparisons"])
check("Each comparison: field, both sources, both values, result, evidence",
      all(set(c) >= {"field", "source_a", "value_a", "source_b", "value_b", "result", "evidence"}
          for c in v["comparisons"]))
low = dict(r["fields"])
low["duty_amount_ghs"] = dict(low["duty_amount_ghs"], confidence="LOW", method="ocr")
ocr = V.validate(doctypes.get(), low, r["hub"], r["request_fields"], identity=r["identity"])
check("A required value read by OCR: NEEDS_REVIEW (LOW_CONFIDENCE_EXTRACTION), with remediation",
      ocr["decision"] == "NEEDS_REVIEW" and ocr["code"] == "LOW_CONFIDENCE_EXTRACTION"
      and ocr["remediation"], ocr["decision"])
check("...confirmed by a person against the PDF: VALID",
      V.validate(doctypes.get(), low, r["hub"], r["request_fields"], identity=r["identity"],
                 ocr_confirmed=True)["decision"] == "VALID")
mixed = V.validate(doctypes.get(), low, dict(r["hub"], bol_awb="176-00000001"),
                   r["request_fields"], identity={"decision": "MATCH"})
check("A mismatch AND a review reason: INVALID (a failure is never softened to review)",
      mixed["decision"] == "INVALID" and mixed["code"] == "CROSS_VALIDATION_FAILED", mixed["code"])

# ═════════════════════════════════════════════════════════════════════════
rule("10. OUTPUT — atomic, never overwritten, verified before any email")
# ═════════════════════════════════════════════════════════════════════════
st = store("out")
r = job(st)
check("No partial file is left in the output folder",
      not list(st.output_dir.glob(".*.partial")), list(st.output_dir.iterdir()))
xs = list(st.output_dir.glob("*.xlsx"))
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(xs[0])
ws = wb[doctypes.get()["template"]["sheet"]]
ws["G19"] = 0
wb.save(xs[0])
r2 = st.get(r["po_id"])
r2["output"]["sha256"] = hashlib.sha256(xs[0].read_bytes()).hexdigest()
check("A template formula replaced after generation is caught before sending",
      any("formula G19" in x for x in P.blocked_reasons(st, r2)), P.blocked_reasons(st, r2))

# ═════════════════════════════════════════════════════════════════════════
rule("11. EMAIL — a distributed transaction (stand-in Graph)")
# ═════════════════════════════════════════════════════════════════════════
st = store("mail")
waits = []
r = job(st)
GRAPH.state["plan"]["create"] = [429]
GRAPH.state["retry_after"] = 7
r, o = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=waits.append)
check("429 on the draft: retried, honouring Retry-After (7 s), then CONFIRMED once",
      o == "CONFIRMED" and 7.0 in waits and len(sends_for(r)) == 1, (o, waits))
r = job(st)
GRAPH.state["plan"]["send"] = [503]
r, o = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=NOSLEEP)
check("503 on the send: reconciled (still a draft) → sent once on the SAME draft → CONFIRMED",
      o == "CONFIRMED" and len(sends_for(r)) == 1, (o, len(sends_for(r))))
r = job(st)
GRAPH.state["plan"]["send"] = ["drop"]
r, o = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=NOSLEEP)
check("Accepted, then the response lost: reconciled in Sent Items → CONFIRMED, sent ONCE",
      o == "CONFIRMED" and len(sends_for(r)) == 1, (o, len(sends_for(r))))
r = job(st)
GRAPH.state["plan"]["send"] = ["mutate"]
r, o = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=NOSLEEP)
check("The sent message is not what was prepared (another recipient): "
      "EMAIL_RECONCILIATION_FAILED — never CONFIRMED, never resent",
      o == "MISMATCH" and r["state"] == S.EMAIL_RECONCILIATION_FAILED
      and r["email"]["reconciliation"]["mismatches"], (o, r["state"]))
again, o2 = P.send(st, st.get(r["po_id"]), M.GraphMailer(), by="omar", sleep=NOSLEEP)
check("...and a Send pressed again is BLOCKED", o2 == "BLOCKED" and len(sends_for(r)) == 1)
r = job(st)
GRAPH.state["plan"]["send"] = [400]
r, o = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=NOSLEEP)
check("A permanent Graph error (400): EMAIL_FAILED (EMAIL_SUBMISSION_FAILED), nothing sent",
      o == "FAILED" and r["failure"]["code"] == "EMAIL_SUBMISSION_FAILED"
      and not sends_for(r), (o, r.get("failure")))
r = job(st)
r, o = P.send(st, r, M.GraphMailer(), by="omar", confirm_wait_s=0, sleep=NOSLEEP)
r2, o2 = P.send(st, st.get(r["po_id"]), M.GraphMailer(), by="omar", sleep=NOSLEEP)
check("A COMPLETED job sent again: BLOCKED — no second email", o == "CONFIRMED" and
      o2 == "BLOCKED" and len(sends_for(r)) == 1)
dup = job(st, ref=r["reference"], data=Path(st.get(r["po_id"])["document"]["evidence"])
          .read_bytes())
check("The same Bill of Entry submitted again as a new job: SKIPPED_DUPLICATE",
      dup["state"] == S.SKIPPED and dup["skip_reason"] == "SKIPPED_DUPLICATE")
check("Graph correlation ids are kept on the job (never the token)",
      st.get(r["po_id"])["email"].get("draft_request_id")
      and "tok-" not in json.dumps(st.get(r["po_id"])))

# ═════════════════════════════════════════════════════════════════════════
rule("12. CRASH CONSISTENCY — a worker killed at every transition, then recovered")
# ═════════════════════════════════════════════════════════════════════════
CHILD = HERE / "fixtures" / "po_crash_child.py"
points = ["transition:{0}".format(n) for n in range(1, 13)] + ["after_create", "after_send"]
results = {}
for point in points:
    st = store("crash-" + point.replace(":", ""))
    ref = fresh_ref()
    pdf = WORK / (ref + ".pdf")
    pdf.write_bytes(F.pdf_of(F.boe_text(bl=ref)))
    rec = st.create(doctypes.DEFAULT, ref, {})
    proc = subprocess.run([sys.executable, str(CHILD), str(st.folder), rec["po_id"], str(pdf),
                           point], cwd=str(HERE), env=dict(os.environ), capture_output=True,
                          text=True, timeout=120)
    crashed = proc.returncode == 137
    # The worker comes back: recovery first, then the normal sweep would send.
    P.recover(st, CONFIG, source=F.Source(pdf.read_bytes(), ref=ref),
              mailer=M.GraphMailer(), log=lambda *a: None)
    r = st.get(rec["po_id"])
    if r["state"] in (S.EMAIL_PREPARED, S.EMAIL_FAILED):
        r, _o = P.send(st, r, M.GraphMailer(), by="recover", confirm_wait_s=0, sleep=NOSLEEP)
    if r["state"] == S.EMAIL_SENT:
        r, _o = P.confirm(st, r, M.GraphMailer(), wait_s=0)
    bad = invariants(st, r)
    results[point] = (crashed, r["state"], len(sends_for(r)), bad)
    check("Killed at {0}: recovered to EMAIL_CONFIRMED, exactly one email, invariants hold"
          .format(point), crashed and r["state"] == S.EMAIL_CONFIRMED and
          len(sends_for(r)) == 1 and not bad,
          (crashed, r["state"], len(sends_for(r)), bad, proc.stderr[-300:]))

# ═════════════════════════════════════════════════════════════════════════
rule("13. CONCURRENCY — processes racing for the same document and the same send")
# ═════════════════════════════════════════════════════════════════════════


RACER = HERE / "fixtures" / "po_race_child.py"


def race(kind, folder, target, pdf="-", n=4):
    """n separate OS processes, released at the same instant; their outcomes."""
    start_at = time.time() + 2.5
    procs = [subprocess.Popen([sys.executable, str(RACER), kind, str(folder), target, str(pdf),
                               str(start_at)], cwd=str(HERE), env=dict(os.environ),
                              stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
             for _ in range(n)]
    out = []
    for p in procs:
        stdout, stderr = p.communicate(timeout=180)
        out.append((stdout.strip().splitlines() or ["CRASHED: " + stderr[-200:]])[-1])
    return sorted(out)


st = store("race")
ref = fresh_ref()
pdf = WORK / (ref + ".pdf")
pdf.write_bytes(F.pdf_of(F.boe_text(bl=ref)))
states = race("process", st.folder, ref, pdf)
check("4 processes, the same Bill of Entry, at the same instant: exactly 1 output, "
      "3 SKIPPED_DUPLICATE", states.count(S.EMAIL_PREPARED) == 1 and
      states.count(S.SKIPPED) == 3 and len(list(st.output_dir.glob("*.xlsx"))) == 1, states)
target = next(r for r in st.all() if r["state"] == S.EMAIL_PREPARED)
outcomes = race("send", st.folder, target["po_id"])
final = st.get(target["po_id"])
check("4 processes press Send on the same job at the same instant: ONE email, the others "
      "BLOCKED", len(sends_for(final)) == 1 and outcomes.count("CONFIRMED") == 1
      and final["state"] == S.EMAIL_CONFIRMED and not invariants(st, final), outcomes)
check("...and the audit chain stayed intact under the concurrent writers",
      st.verify_audit()["ok"], st.verify_audit())

# ═════════════════════════════════════════════════════════════════════════
rule("14. PROPERTIES — random documents and Graph behaviour, invariants every time")
# ═════════════════════════════════════════════════════════════════════════
rng = random.Random(20261006)
st = store("prop")
violations = []
outcomes = {}
for i in range(60):
    ref = fresh_ref()
    text = {"bl": ref}
    pick = rng.choice(["good", "good", "good", "no_inv", "two_inv", "bad_inv", "bl_other",
                       "bad_amount", "dup_vat", "arith", "neg"])
    if pick == "no_inv":
        text["invoice"] = None
    elif pick == "two_inv":
        text["extra"] = "Commercial Invoice Number: B{0}\n".format(i)
    elif pick == "bad_inv":
        text["invoice"] = "26 {0}".format(i)
    elif pick == "bl_other":
        text["bl"] = "176-55555555"
    elif pick == "bad_amount":
        text["duty"] = "653,49{0}.3O".format(i % 10)
    elif pick == "dup_vat":
        text["vat"] = "Import VAT  1,000.00\nImport VAT  1,000.00\n"
    elif pick == "arith":
        text["import_duty"] = "200,000.00"
    elif pick == "neg":
        text["duty"] = "-653,492.35"
    request = rng.choice([{}, {}, {"invoice_no": "9116093"}, {"invoice_no": "1"}])
    r = job(st, ref=ref, data=F.pdf_of(F.boe_text(**text)), request=request)
    plan = rng.choice([[], [], [503], ["drop"], [400], [429]])
    GRAPH.state["plan"]["send"] = list(plan)
    GRAPH.state["retry_after"] = 1
    if r["state"] == S.EMAIL_PREPARED:
        r, _o = P.send(st, r, M.GraphMailer(), by="prop", confirm_wait_s=0, sleep=NOSLEEP)
        if r["state"] == S.EMAIL_UNKNOWN:
            r, _o = P.reconcile(st, r, M.GraphMailer(), wait_s=0)
        # Pressing Send again must never produce a second email.
        P.send(st, st.get(r["po_id"]), M.GraphMailer(), by="prop", confirm_wait_s=0,
               sleep=NOSLEEP)
        r = st.get(r["po_id"])
    GRAPH.state["plan"]["send"] = []
    outcomes[r["state"]] = outcomes.get(r["state"], 0) + 1
    bad = invariants(st, r)
    if pick != "good" and r["state"] in (S.EMAIL_PREPARED, S.EMAIL_CONFIRMED, S.EMAIL_SENT):
        bad.append("a {0} document reached {1}".format(pick, r["state"]))
    if bad:
        violations.append((i, pick, request, plan, r["state"], bad))
print("    outcomes over 60 random jobs: {0}".format(dict(sorted(outcomes.items()))))
check("60 random jobs: every invariant holds; no bad document is ever generated or sent; never "
      "two emails for one document", not violations, violations[:3])
check("...and the random mix really exercised the paths (confirmed, review, failed, mismatch)",
      outcomes.get(S.EMAIL_CONFIRMED, 0) > 5 and outcomes.get(S.NEEDS_REVIEW, 0) > 0 and
      (outcomes.get(S.VALIDATION_FAILED, 0) + outcomes.get(S.EXTRACTION_FAILED, 0)) > 0, outcomes)

# ═════════════════════════════════════════════════════════════════════════
rule("15. OBSERVABILITY — metrics from the persisted jobs only")
# ═════════════════════════════════════════════════════════════════════════
m = Q.metrics(st)["not_real"]
check("Queue depth, active jobs, NEEDS_REVIEW, failure / unknown / retry rates, duplicates — "
      "all present and computed from records",
      all(k in m for k in ("queue_depth", "active", "needs_review", "pdf_failure",
                           "extraction_failure", "validation_failure", "email_failure",
                           "email_unknown", "retry", "duplicates_prevented", "recovery_success",
                           "avg_processing_ms", "avg_stage_ms"))
      and m["needs_review"] == outcomes.get(S.NEEDS_REVIEW, 0), m)
check("A rate with nothing to divide by is null, never invented",
      Q.metrics(store("empty"))["real"]["email_failure"]["pct"] is None)

GRAPH.stop()
shutil.rmtree(WORK, ignore_errors=True)
print()
print("{0} passed, {1} failed".format(len(PASS), len(FAIL)))
sys.exit(1 if FAIL else 0)
