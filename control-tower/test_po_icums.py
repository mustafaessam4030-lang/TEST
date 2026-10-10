"""
PO Automation on a REAL ICUMS Bill of Entry layout — the business's own
example: the "Sheet4" BOE of the Ghana Duty Payment cheque-request workbook
(BOE 40726534505 / 00, BL/AWB J552493) and its completed "Duty Template".

    python test_po_icums.py

The PDF is a replica (fixtures/icums_boe.py) of that BOE: the same boxed
form, the same positions, the same figures, and a text layer that is not in
reading order — like the real one. It is NOT the real eHub download; the
real document is verified on the Windows worker (python -m po read <pdf>).

What it proves:
  1. the defect: the line rules read the real form wrongly (the failure seen
     on the first real BOE — only the user reference, and BL/AWB = box 6's
     reference number);
  2. every highlighted field of the example is read by position;
  3. the saved Duty Template holds exactly the business's figures (G4, G6,
     G8, C13, C19, C21, G20 and the computed G19);
  4. what must stop still stops (totals that do not add up, a malformed
     amount, another shipment's BOE, a non-USD invoice, no G4 source).
"""

import os
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE / "fixtures"))
WORK = Path(tempfile.mkdtemp(prefix="po_icums_"))
os.environ["PO_DATA_DIR"] = str(WORK / "default")
for name in ("PO_DEFAULT_SUPPLIER", "PO_DEFAULT_BRANCH", "PO_DEFAULT_CHARGE_TO",
             "PO_DEFAULT_PRIORITY"):
    os.environ.pop(name, None)

from openpyxl import load_workbook  # noqa: E402

import po_engine as F  # noqa: E402
from icums_boe import TAX_LINES, icums_boe_pdf  # noqa: E402
from po import doctypes, extract as X, pipeline as P, store as S  # noqa: E402

PASS, FAIL = [], []
DOCTYPE = doctypes.get()
# The pilot: every job stops at SAVED, no email.
CONFIG = {"recipient": None, "sender": None, "auto_send": False,
          "defaults": {"supplier": None, "branch": None, "charge_to": None, "priority": None},
          "confirm_wait_s": 0, "email_disabled": True}

# The business's completed Duty Template for this BOE (the example workbook).
EXAMPLE = {"G4": 9116093, "G6": 653670.54, "C13": "40726534505 / 00", "G10": "CAT",
           "G13": "32600.CPA.G005",
           "G20": "=304446.44+1066.03+50741.08+50741.08+177.67+177.67", "G19": "246320.57"}


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print("  {0}  {1}{2}".format("PASS" if ok else "FAIL", name,
                                 "  ({0})".format(str(detail)[:500]) if detail and not ok else ""))


def rule(title):
    print("\n" + "=" * 74 + "\n" + title + "\n" + "=" * 74)


def read(data):
    r = X.read_pdf(data)
    return r, X.extract(r["text"], DOCTYPE, r["spans"], r["words"])


_n = {"i": 0}


def job(data, ref="J552493", identifier="40726534505", request=None):
    _n["i"] += 1
    st = S.Store(folder=WORK / "job{0}".format(_n["i"]))
    record = st.create(doctypes.DEFAULT, ref, request or {})
    return P.process(st, record, F.Source(data, ref=ref, identifier=identifier), CONFIG,
                     sleep=lambda s: None)


def value(fields, name):
    f = fields.get(name) or {}
    return f.get("value") if f.get("status") == X.FOUND else None


EXAMPLE_PDF = icums_boe_pdf()

# ═════════════════════════════════════════════════════════════════════════
rule("1. THE DEFECT: THE LINE RULES CANNOT READ THE BOXED FORM")
# ═════════════════════════════════════════════════════════════════════════
r = X.read_pdf(EXAMPLE_PDF)
old = X.extract(r["text"], DOCTYPE, r["spans"])          # text only, as before the fix
found = sorted(n for n, f in old.items() if f["status"] == X.FOUND)
check("Read as text lines (the old way), the example BOE gives the real failure: the user "
      "reference found, the declaration, date, amounts missing — not a Bill of Entry",
      value(old, "user_reference") == "DDAO9116093" and not X.looks_like(old)
      and all(old[n]["status"] == X.MISSING for n in ("document_number", "document_date",
                                                      "cif_usd", "exchange_rate",
                                                      "duty_amount_ghs")), found)
check("...and BL/AWB takes box 6's reference number — the wrong value the real run showed",
      value(old, "bl_awb") == "2607150575GCH000154", value(old, "bl_awb"))
check("The positioned words identify the ICUMS form", X.is_icums_form(r["words"]))
check("A plain text document is not taken for the ICUMS form (the line rules still read it)",
      not X.is_icums_form(X.read_pdf(F.pdf_of(F.boe_text()))["words"]))

# ═════════════════════════════════════════════════════════════════════════
rule("2. EVERY HIGHLIGHTED FIELD, READ WHERE THE FORM PRINTS IT")
# ═════════════════════════════════════════════════════════════════════════
r, f = read(EXAMPLE_PDF)
expect = {"bl_awb": "J552493", "document_number": "40726534505 / 00",
          "document_date": "2026-07-15", "user_reference": "DDAO9116093",
          "delivery_terms": "CPT ACCRA", "cif_usd": 174071.09, "invoice_currency": "USD",
          "exchange_rate": 11.4857, "rate_currency": "USD", "fob_ncy": 1776621.53,
          "duty_amount_ghs": 653670.54, "stated_import_duty": 163253.97,
          "doc_status": "Assessed"}
for name, want in expect.items():
    check("{0} = {1!r}".format(name, want), value(f, name) == want,
          (f.get(name) or {}).get("status"), )
check("Each value keeps its evidence: the label, the value and where it was printed",
      f["bl_awb"]["evidence"].startswith("BL/AWB") and f["bl_awb"]["page"] == 1 and
      f["cif_usd"]["evidence"].startswith("Total Invoice Fcy") and
      f["cif_usd"]["confidence"] == "HIGH", (f["bl_awb"]["evidence"], f["cif_usd"]["evidence"]))
vat = value(f, "vat_lines") or []
check("The six VAT / levy lines of G20, in the Duty Template's order, with their tax codes",
      [(l["label"], l["code"], l["amount"]) for l in vat] == [
          ("Import VAT", "02", 304446.44), ("Network Charge VAT", "33", 1066.03),
          ("Import NHIL", "47", 50741.08), ("GETFund Import", "88", 50741.08),
          ("Network Charge NHIL", "48", 177.67), ("Network Charge GET Fund Levy", "89", 177.67)],
      [(l["label"], l["code"], l["amount"]) for l in vat])
tax = value(f, "tax_lines") or []
check("All 18 lines of B ACCOUNTING DETAILS, names and codes as printed",
      [(l["label"], l["code"]) for l in tax] == [(n, c) for n, c, _e, _p in TAX_LINES],
      [(l["label"], l["code"]) for l in tax])
check("...amounts from the Amount Payable column — never box 40's per-item figures beside them",
      [l["amount"] for l in tax] == [X.to_float(p) for _n, _c, _e, p in TAX_LINES]
      and 1079.63 not in [l["amount"] for l in tax])
check("...and they add up to the printed Total, to the cent",
      round(sum(l["amount"] for l in tax), 2) == 653670.54)

for kw, what in (({"reading_order": True}, "the text layer in reading order"),
                 ({"font_scale": 0.85}, "smaller print"),
                 ({"font_scale": 0.9, "reading_order": True}, "smaller print, reading order")):
    _r, g = read(icums_boe_pdf(**kw))
    check("The same values from {0}".format(what),
          all(value(g, n) == w for n, w in expect.items()) and
          [l["amount"] for l in value(g, "vat_lines") or []] == [l["amount"] for l in vat] and
          len(value(g, "tax_lines") or []) == 18,
          {n: (g[n]["status"], g[n].get("note")) for n in g if g[n]["status"] != X.FOUND})

# ═════════════════════════════════════════════════════════════════════════
rule("3. THE PIPELINE: THE BUSINESS'S DUTY TEMPLATE, CELL FOR CELL (NO EMAIL)")
# ═════════════════════════════════════════════════════════════════════════
rec = job(EXAMPLE_PDF)
check("The job runs end to end and stops at SAVED (pilot: no email prepared or sent)",
      rec["state"] == S.SAVED and rec["validation"]["decision"] == "VALID"
      and not rec.get("email", {}).get("message_id"), (rec["state"], rec.get("failure"),
                                                       rec.get("validation", {}).get("reasons")))
g4 = rec["request_fields"]["invoice_no"]
check("G4 = 9116093: the digits of the Bill of Entry's own User Reference DDAO9116093 (it "
      "prints no Invoice No.), with that origin and evidence recorded",
      g4["value"] == "9116093" and g4["origin"] == "bill_of_entry:user_reference"
      and "DDAO9116093" in g4["evidence"], g4)
calc = {c["rule"]: c for c in rec["calculations"]}
check("Calculations: G20 = 407,349.97; G19 = G6 − G20 = 246,320.57 (as the example); "
      "Σ tax lines = Total → OK",
      calc["G20_VAT_BLOCK"]["result"] == "407349.97" and
      calc["G19_IMPORT_DUTY"]["result"] == EXAMPLE["G19"] and
      calc["XCHECK_TAX_TOTAL"]["result"] == "OK" and "XCHECK_IMPORT_DUTY" not in calc, calc)
check("C24 = G6 ÷ C21 and C26 = C24 ÷ C19 from the BOE's own rate and invoice value",
      calc["C24_DUTY_USD"]["result"] == "56911.68" and calc["C26_DUTY_SHARE"]["result"] ==
      "0.3269", (calc["C24_DUTY_USD"]["result"], calc["C26_DUTY_SHARE"]["result"]))
ws = load_workbook(rec["output"]["path"])["Duty Template"]
for cell in ("G4", "G6", "C13", "G10", "G13", "G20"):
    check("Saved workbook {0} = the example's {1!r}".format(cell, EXAMPLE[cell]),
          ws[cell].value == EXAMPLE[cell], ws[cell].value)
check("C19 = Total Invoice Fcy 174,071.09 (USD) and C21 = Rate of Xchange 11.4857, from the BOE",
      ws["C19"].value == 174071.09 and ws["C21"].value == 11.4857,
      (ws["C19"].value, ws["C21"].value))
check("G8 = the BOE's Date 15/07/2026", ws["G8"].value.date().isoformat() == "2026-07-15",
      ws["G8"].value)
check("The template's own formulas are kept (G19, C24, C26, G41, F16), never overwritten",
      all(ws[c].value == fm for c, fm in DOCTYPE["template"]["formulas"].items()))
check("Branch (G11) and priority (C9) are not on the BOE and not configured: left empty",
      ws["G11"].value is None and ws["C9"].value is None, (ws["G11"].value, ws["C9"].value))

# ═════════════════════════════════════════════════════════════════════════
rule("4. WHAT MUST STOP STILL STOPS")
# ═════════════════════════════════════════════════════════════════════════
r = job(icums_boe_pdf(total="653,700.54"))
check("The tax lines do not add up to the printed Total (off by 30.00): VALIDATION_FAILED, "
      "BUSINESS_RULE_FAILED, nothing saved",
      r["state"] == S.VALIDATION_FAILED and r["validation"]["code"] == "BUSINESS_RULE_FAILED"
      and not r.get("output"), (r["state"], r["validation"].get("reasons")))
bad = [list(t) for t in TAX_LINES]
bad[1][3] = "3O4,446.44"
r = job(icums_boe_pdf(tax_lines=[tuple(t) for t in bad]))
check("A malformed amount (3O4,446.44 — a letter O): NORMALIZATION_FAILED, never coerced",
      r["state"] == S.EXTRACTION_FAILED and r["failure"]["code"] == "NORMALIZATION_FAILED"
      and not r.get("output"), (r["state"], r.get("failure")))
r = job(icums_boe_pdf(bl="K998877"))
check("Another shipment's BOE (BL/AWB K998877 on record J552493): stopped, both values shown",
      r["state"] == S.IDENTITY_MISMATCH and not r.get("output") and
      any("K998877" in x and "J552493" in x for x in r["validation"]["reasons"]),
      (r["state"], r["validation"].get("reasons")))
r = job(icums_boe_pdf(bl=""))
check("No BL/AWB printed: MISSING — box 6's reference number is never taken instead",
      r["fields"]["bl_awb"]["status"] == X.MISSING and r["state"] == S.VALIDATION_FAILED,
      r["fields"]["bl_awb"])
r = job(icums_boe_pdf(invoice_currency="EUR", rate_currency="EUR"))
check("An invoice in EUR (C19 is 'Amount (USD)'): BUSINESS_RULE_FAILED, nothing saved",
      r["state"] == S.VALIDATION_FAILED and r["validation"]["code"] == "BUSINESS_RULE_FAILED"
      and not r.get("output"), r["validation"].get("reasons"))
r = job(icums_boe_pdf(rate_currency="GBP"))
check("A rate of exchange for another currency than the invoice's: stopped",
      r["state"] == S.VALIDATION_FAILED and
      any("Exchange-rate currency" in x for x in r["validation"]["reasons"]),
      r["validation"].get("reasons"))
r = job(icums_boe_pdf(user_reference=None))
check("No Invoice No. and no User Reference: G4 has no source → NEEDS_REVIEW, nothing saved",
      r["state"] == S.NEEDS_REVIEW and r["validation"]["code"] == "G4_SOURCE_UNPROVEN"
      and not r.get("output"), (r["state"], r["validation"].get("code")))
r = job(icums_boe_pdf(user_reference="DD-AO/91-16"))
check("A User Reference that is not letters-then-digits is not turned into a number: review",
      r["state"] == S.NEEDS_REVIEW and r["validation"]["code"] == "G4_SOURCE_UNPROVEN",
      (r["state"], r["request_fields"]["invoice_no"]))
r = job(EXAMPLE_PDF, request={"invoice_no": "7001234"})
check("The job says invoice 7001234, the BOE's User Reference gives 9116093: conflict → review",
      r["state"] == S.NEEDS_REVIEW and r["validation"]["code"] == "G4_SOURCE_UNPROVEN",
      (r["state"], r["validation"].get("reasons")))
r = job(icums_boe_pdf(extra_footer="Invoice No: 7001234"))
check("An explicit 'Invoice No.' printed on the BOE wins over the User Reference",
      r["request_fields"]["invoice_no"]["value"] == "7001234" and
      r["request_fields"]["invoice_no"]["origin"] == "bill_of_entry",
      r["request_fields"]["invoice_no"])
codes = [list(t) for t in TAX_LINES]
codes[1][1] = "03"
r = job(icums_boe_pdf(tax_lines=[tuple(t) for t in codes]))
check("A VAT line printed with another tax code (Import VAT under 03, not 02): the block is not "
      "trusted — stopped, nothing saved",
      r["fields"]["vat_lines"]["status"] == X.AMBIGUOUS and r["state"] == S.VALIDATION_FAILED
      and not r.get("output"), (r["fields"]["vat_lines"].get("note"), r["state"]))
r = job(icums_boe_pdf(font_scale=1.25))
check("Names printed over the Code column (an overprinted form): NORMALIZATION_FAILED — never a "
      "VAT line quietly dropped from G20",
      r["state"] == S.EXTRACTION_FAILED and r["failure"]["code"] == "NORMALIZATION_FAILED"
      and not r.get("output"), (r["state"], r.get("failure")))
r, f = read(icums_boe_pdf(date="15/07/2026", total=None))
check("No Total row printed: the total duty is MISSING, not summed by us",
      f["duty_amount_ghs"]["status"] == X.MISSING, f["duty_amount_ghs"])

# ═════════════════════════════════════════════════════════════════════════
rule("7. TWO PRINTED DECLARATION NUMBERS (job 41026844263, K284375): AMBIGUOUS, "
     "SETTLED ONLY BY EVIDENCE")
# ═════════════════════════════════════════════════════════════════════════
# A sanitised replica: page 2 repeats the header's "Bill of Entry(BOE) No :"
# label. No real document content is used.
from po.__main__ import trace_lines  # noqa: E402
OTHER = icums_boe_pdf(continuation_number="40726599999 / 00")
r, fields = read(OTHER)
dn = fields["document_number"]
check("Reproduced: the label printed on two pages with two numbers is AMBIGUOUS, "
      "each candidate with its page",
      dn["status"] == X.AMBIGUOUS and sorted((c["value"], c["page"]) for c in dn["candidates"])
      == [("40726534505 / 00", 1), ("40726599999 / 00", 2)], dn.get("candidates"))
check("...it is still read as an ICUMS Bill of Entry, by position (the right schema)",
      X.is_icums_form(r["words"]) and X.looks_like(fields))
same = read(icums_boe_pdf(continuation_number="40726534505 / 00"))[1]["document_number"]
check("The same number printed again on page 2 is one value, not an ambiguity",
      same["status"] == X.FOUND and same["value"] == "40726534505 / 00")
spaced = read(icums_boe_pdf(continuation_number="40726534505/00"))[1]["document_number"]
check("...nor is the same number with different spacing around '/'",
      spaced["status"] == X.FOUND and spaced["value"] == "40726534505 / 00")
j = job(OTHER, identifier="40726534505")
f = j["fields"]["document_number"]
check("eHub's own number for the document matches exactly one printed value: that "
      "printed value is used",
      f["status"] == X.FOUND and f["value"] == "40726534505 / 00" and f["page"] == 1, f)
check("...with the rule, eHub's number and the rejected candidate kept as evidence",
      f["resolution"]["identifier"] == "40726534505" and
      [c["value"] for c in f["resolution"]["rejected"]] == ["40726599999 / 00"] and
      f["resolution"]["rejected"][0]["page"] == 2)
check("...and only then do validation and the template run, to the same end as the "
      "one-page document", j["state"] == job(EXAMPLE_PDF)["state"] and j.get("output"))
events = [e for e in S.Store(folder=WORK / "job{0}".format(_n["i"] - 1)).events(j["po_id"])
          if e["event"] == "FIELDS_EXTRACTED"]
check("...and the extraction event says the field was settled",
      events and events[-1]["metadata"].get("resolved") == ["document_number"],
      events[-1]["metadata"] if events else None)
for number, ident, why in (("40726534505 / 01", "40726534505", "same number, other suffix"),
                           ("40726599999 / 00", "40726599990", "no printed value is eHub's")):
    k = job(icums_boe_pdf(continuation_number=number), identifier=ident)
    g = k["fields"]["document_number"]
    req = [c for c in (k.get("validation") or {}).get("checks", [])
           if c["name"] == "required:document_number"]
    check("Not settled ({0}): stays AMBIGUOUS, says why, and nothing is chosen".format(why),
          g["status"] == X.AMBIGUOUS and g.get("value") is None and g.get("note"), g.get("note"))
    check("...validation blocks it, and no template is generated ({0})".format(why),
          k["state"] in (S.VALIDATION_FAILED, S.NEEDS_REVIEW) and not k.get("output") and
          req and req[0]["status"] == "AMBIGUOUS" and req[0]["blocking"],
          (k["state"], req))
    if why.startswith("same"):
        lines = "\n".join(trace_lines(k))
        check("'python -m po trace' shows the blocking field and both printed candidates "
              "with their pages", "document_number (Declaration (BOE) No.): AMBIGUOUS" in lines
              and "40726534505 / 00 (page 1" in lines and "40726534505 / 01 (page 2" in lines,
              lines)
from po.__main__ import find_job  # noqa: E402
empty_folder, job_folder = WORK / "trace_empty", WORK / "job{0}".format(_n["i"])
for key in (k["identifier"], k["reference"], k["po_id"]):
    hit, where, searched = find_job(key, [empty_folder, job_folder])
    check("'trace {0}' finds the job in whichever folder holds it".format(key),
          hit is not None and hit["po_id"] == k["po_id"] and where == job_folder, searched)
miss, _w, searched = find_job("99999999999", [empty_folder, job_folder])
check("...and when there is none, says which folders it searched and how many jobs each held",
      miss is None and [n for _p, n in searched] == [0, 1], searched)
check("Without an eHub number nothing is settled",
      X.resolve_by_identifier(dn, None)["status"] == X.AMBIGUOUS and
      X.resolve_by_identifier(dn, "")["status"] == X.AMBIGUOUS)
check("A value never printed in the document is never used, even if eHub has it",
      X.resolve_by_identifier(dn, "12345678901")["status"] == X.AMBIGUOUS)

# ═════════════════════════════════════════════════════════════════════════
rule("8. THE REAL DOCUMENT OF JOB 41026844263: '/ 00' AND A BARE NUMBER, A STAMP OVER THE "
     "TAX TABLE, A HEADING ON TWO ROWS")
# ═════════════════════════════════════════════════════════════════════════
# What the trace of the real document showed (sanitised replica, no real
# content): page 1 prints "41026844263 / 00", page 2 "41026844263"; the
# clearing stamp "Validated Entry / accepted by Declarant" falls across the
# EXIM and Inspection Fee rows; "Exempted/Suspended" sits on its own row.
from po.__main__ import reread_lines  # noqa: E402
bare = icums_boe_pdf(continuation_number="40726534505")
dn2 = read(bare)[1]["document_number"]
check("Reproduced: '/ 00' on page 1 and the bare number on page 2 read as two values (AMBIGUOUS)",
      dn2["status"] == X.AMBIGUOUS and sorted(c["value"] for c in dn2["candidates"]) ==
      ["40726534505", "40726534505 / 00"], dn2.get("candidates"))
j = job(bare, identifier="40726534505")
f = j["fields"]["document_number"]
kept = S.Store(folder=WORK / "job{0}".format(_n["i"])).get(j["po_id"])
check("The saved job keeps the full evidence of the other printing (value, page, line)",
      kept["fields"]["document_number"]["resolution"]["rejected"] ==
      [{"value": "40726534505", "page": 2, "line": f["resolution"]["rejected"][0]["line"]}]
      and f["resolution"]["rejected"][0]["line"],
      kept["fields"]["document_number"].get("resolution"))
check("With eHub's number: one declaration — the full printed form '/ 00' (page 1) is used",
      f["status"] == X.FOUND and f["value"] == "40726534505 / 00" and f["page"] == 1, f)
check("...the rule says why, and the bare printing is kept as evidence",
      "suffix / 00" in f["resolution"]["rule"] and
      [(c["value"], c["page"]) for c in f["resolution"]["rejected"]] == [("40726534505", 2)],
      f.get("resolution"))
check("...and the job reaches the same end as the one-page document, C13 = '40726534505 / 00'",
      j["state"] == job(EXAMPLE_PDF)["state"] and j.get("output") and
      j["fields"]["document_number"]["value"] == EXAMPLE["C13"], j["state"])
check("Without eHub's number the two printings stay AMBIGUOUS (nothing is assumed)",
      X.resolve_by_identifier(dn2, None)["status"] == X.AMBIGUOUS)
three = dict(dn2, candidates=dn2["candidates"] + [dict(dn2["candidates"][0],
                                                       value="40726534505 / 01", page=3)])
check("'/ 00', '/ 01' and the bare number: two different suffixes stay AMBIGUOUS",
      X.resolve_by_identifier(three, "40726534505")["status"] == X.AMBIGUOUS)

RED, BLUE = (0.8, 0.0, 0.0), (0.0, 0.0, 0.7)
r, f = read(icums_boe_pdf(stamp={"size": 6.5}))
t = f["tax_lines"]
check("Reproduced: a stamp printed EXACTLY like the table makes the rows MALFORMED — the same "
      "rows as the real trace, and nothing is guessed",
      t["status"] == X.MALFORMED and "(EXIM) Validated Levy 87 Entry 0.00" in t["note"],
      t.get("note"))
k = job(icums_boe_pdf(stamp={"size": 6.5}))
check("...the job stops at EXTRACTION_FAILED (NORMALIZATION_FAILED), no template",
      k["state"] == S.EXTRACTION_FAILED and k["failure"]["code"] == "NORMALIZATION_FAILED"
      and not k.get("output"), (k["state"], k.get("failure")))
folder = WORK / "job{0}".format(_n["i"])
before = S.Store(folder=folder).get(k["po_id"])
again = "\n".join(reread_lines(k, folder))
check("'trace' re-reads the kept PDF with this version and shows how each table word is "
      "printed (size, colour, font)",
      "RE-READ NOW" in again and "tax_lines" in again and "Validated[6.5pt #000000" in again
      and "the table's figures are printed 6.5 pt #000000" in again, again[-600:])
after = S.Store(folder=folder).get(k["po_id"])
check("...and the re-read changes nothing: same state, same fields, same events",
      after["state"] == before["state"] and after["fields"] == before["fields"] and
      len(S.Store(folder=folder).events(k["po_id"])) ==
      len(S.Store(folder=folder).events(k["po_id"])))
for style, why in (({"size": 6.5, "color": RED}, "another colour"),
                   ({"size": 9.0}, "a larger size")):
    r, f = read(icums_boe_pdf(stamp=style))
    t = f["tax_lines"]
    want = [(n, c, X.to_float(a)) for n, c, _e, a in TAX_LINES]
    got = [(l["label"], l["code"], l["amount"]) for l in t.get("value") or []]
    check("A stamp printed in {0}: its words set aside, every tax line read as printed".format(
          why), t["status"] == X.FOUND and got == want, (t.get("note"), got[13:17]))
    check("...the words set aside are kept as evidence, with why ({0})".format(why),
          [w["text"] for w in t["set_aside"]][:1] == ["Validated"] and
          all(w["why"] and w["line"] for w in t["set_aside"]) and "stamp" in t["note"],
          t.get("set_aside"))
s = job(icums_boe_pdf(stamp={"size": 6.5, "color": RED}))
check("...and the job reaches the same end as the unstamped document (Σ lines = Total)",
      s["state"] == job(EXAMPLE_PDF)["state"] and s.get("output"), s["state"])
check("...and 'trace' lists the stamp words it set aside, with why",
      "set aside (stamp): 'Validated' — printed in another colour" in "\n".join(trace_lines(s)))
digit = (10, 10, 30, 16, "12/09/2026", {"size": 9, "color": 0xCC0000, "upright": True})
ref = {"size": 6.5, "color": 0}
check("A word with a digit is never set aside, however it is printed (no figure is dropped)",
      X._overprint(digit, ref) is None)
check("A word printed at an angle is set aside; one printed like the table is not",
      X._overprint((0, 0, 5, 30, "Validated", {"size": 6.5, "color": 0, "upright": False}),
                   ref) == "printed at an angle" and
      X._overprint((0, 0, 30, 6, "Levy", {"size": 6.5, "color": 0, "upright": True}),
                   ref) is None)
check("A word with no print information (another reader) is never set aside",
      X._overprint((0, 0, 30, 6, "Validated"), ref) is None)
f = read(icums_boe_pdf(split_heading=True))[1]
check("'Exempted/Suspended' on its own row under the heading is the heading, not a tax line",
      f["tax_lines"]["status"] == X.FOUND and len(f["tax_lines"]["value"]) == len(TAX_LINES),
      f["tax_lines"].get("note"))
real = job(icums_boe_pdf(continuation_number="40726534505", split_heading=True,
                         stamp={"size": 6.5, "color": BLUE}), identifier="40726534505")
check("All three at once (the real document's shape): the job reaches the same end as the "
      "clean document, C13 '40726534505 / 00', every value verified by validation",
      real["state"] == job(EXAMPLE_PDF)["state"] and real.get("output") and
      real["fields"]["document_number"]["value"] == "40726534505 / 00" and
      real["validation"]["passed"], (real["state"], real.get("failure")))
real_bad = job(icums_boe_pdf(continuation_number="40726534505 / 01", split_heading=True,
                             stamp={"size": 6.5, "color": BLUE}), identifier="40726534505")
check("...but with two different suffixes printed, it still stops for review — no template",
      real_bad["state"] in (S.VALIDATION_FAILED, S.NEEDS_REVIEW) and not real_bad.get("output"),
      real_bad["state"])

print()
print("{0} passed, {1} failed".format(len(PASS), len(FAIL)))
sys.exit(1 if FAIL else 0)
