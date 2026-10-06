"""
Derive the approved template a document type fills, from the original
workbook — once, reproducibly, with its provenance recorded.

    python -m po.templates.build_template            # DUTY_REQUEST_V1

What it does to the original, and nothing else:

  * keeps only the sheets the type names (for DUTY_REQUEST_V1: "Duty
    Template" and "BOE Template Capture"); the original's other sheets are
    empty working copies — one alone holds ~96,000 formatted blank rows and
    makes every load take 16 seconds. The earlier BOE automation dropped the
    same sheets from its output.
  * clears the input cells, so the approved template carries no earlier
    request's figures.

Layout, formatting, fonts, the logo, merged cells and every formula stay as
the original has them. The manifest records both files' SHA-256, so any
change to either is visible.
"""

import hashlib
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent))

from po import doctypes  # noqa: E402


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def build(doctype_id=None):
    from openpyxl import load_workbook
    doctype = doctypes.get(doctype_id)
    spec = doctype["template"]
    source = HERE / spec["source_file"]
    target = HERE / spec["file"]
    wb = load_workbook(source)
    for name in list(wb.sheetnames):
        if name not in spec["keep_sheets"]:
            del wb[name]
    ws = wb[spec["sheet"]]
    for ref in spec["input_cells"]:
        ws[ref] = None
    wb.active = wb.sheetnames.index(spec["sheet"])
    wb.save(target)
    manifest = {
        "version": spec["version"], "doctype": doctype["id"],
        "file": spec["file"], "sha256": sha256(target),
        "source_file": spec["source_file"], "source_sha256": sha256(source),
        "sheets": list(spec["keep_sheets"]), "cleared_cells": list(spec["input_cells"]),
        "formulas": spec["formulas"],
    }
    (HERE / (spec["version"] + ".manifest.json")).write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    print(json.dumps(build(sys.argv[1] if len(sys.argv) > 1 else None), indent=2))
